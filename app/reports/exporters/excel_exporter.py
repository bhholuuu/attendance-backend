"""Excel exporter (V2.2) using openpyxl.

Produces a small, readable workbook:
  * Summary sheet  - title, meta block, key figures, generated timestamp.
  * Details sheet  - the detailed table (student rollup / daily records / trend).

Everything is written to an in-memory buffer (bytes); no temp files are needed.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .base import build_payload
from .filename import build_filename

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(bold=True)


def export_excel(report_type: str, data: dict) -> bytes:
    """Serialize a report dict into an .xlsx workbook in memory."""
    payload = build_payload(report_type, data)

    wb = Workbook()
    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    ws.append([payload["title"]])
    ws["A1"].font = _TITLE_FONT
    ws.append([payload["subtitle"]])
    ws["A2"].font = _LABEL_FONT

    row = 4
    for label, value in payload["meta"]:
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Summary").font = _LABEL_FONT
    row += 1
    for label, value in payload["summary"]:
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Generated").font = _LABEL_FONT
    ws.cell(
        row=row,
        column=2,
        value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    # ---- Details sheet ----
    details = wb.create_sheet("Details")
    details.append(payload["columns"])
    for col_idx, _ in enumerate(payload["columns"], start=1):
        cell = details.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for record in payload["rows"]:
        details.append(record)

    for col_idx, header in enumerate(payload["columns"], start=1):
        width = min(max(len(str(header)) + 4, 12), 40)
        details.column_dimensions[get_column_letter(col_idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_filename(report_type: str, data: dict) -> str:
    """Safe Excel filename for the given report data."""
    return build_filename(report_type, data, ".xlsx")
